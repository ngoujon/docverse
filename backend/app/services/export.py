import io
import json
from datetime import datetime

from docx import Document as DocxDocument
from docx.shared import Pt
from fpdf import FPDF, XPos, YPos

from .. import models_db

_DEJAVU_REGULAR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def _line(pdf: FPDF, height: float, text: str) -> None:
    """multi_cell() wrapper that always resets the cursor to the left
    margin on the next line afterwards - without this, fpdf2 can leave x
    wherever the last line of text ended, and the following multi_cell()
    then has near-zero width to work with and raises FPDFException."""
    pdf.multi_cell(0, height, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _source_line(source: dict, index: int) -> str:
    if source.get("type") == "web":
        return f"[{index}] {source.get('label', 'web')} - {source.get('url', '')}"
    return f"[{index}] {source.get('label', 'document')}"


def _load_sources(message: models_db.Message) -> list[dict]:
    try:
        return json.loads(message.sources_json or "[]")
    except (json.JSONDecodeError, TypeError):
        return []


def build_pdf(conversation: models_db.Conversation, space_name: str) -> bytes:
    pdf = FPDF()
    pdf.add_font("DejaVu", "", _DEJAVU_REGULAR)
    pdf.add_font("DejaVu", "B", _DEJAVU_BOLD)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    pdf.set_font("DejaVu", "B", 16)
    _line(pdf, 10, conversation.title or "Conversation")

    pdf.set_font("DejaVu", "", 9)
    pdf.set_text_color(110, 110, 110)
    meta = f"Espace : {space_name}  -  Exporte le {datetime.utcnow().strftime('%d/%m/%Y a %H:%M')} UTC"
    _line(pdf, 6, meta)
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    for message in conversation.messages:
        if message.role not in ("user", "assistant"):
            continue
        label = "Question" if message.role == "user" else "Reponse"
        pdf.set_font("DejaVu", "B", 11)
        if message.role == "user":
            pdf.set_text_color(139, 47, 214)
        else:
            pdf.set_text_color(8, 145, 168)
        _line(pdf, 8, label)

        pdf.set_font("DejaVu", "", 11)
        pdf.set_text_color(20, 20, 20)
        _line(pdf, 6.5, message.content)

        sources = _load_sources(message)
        if sources:
            pdf.set_font("DejaVu", "", 9)
            pdf.set_text_color(110, 110, 110)
            for i, source in enumerate(sources, start=1):
                _line(pdf, 5.5, _source_line(source, i))
            pdf.set_text_color(0, 0, 0)
        pdf.ln(5)

    return bytes(pdf.output())


def build_docx(conversation: models_db.Conversation, space_name: str) -> bytes:
    doc = DocxDocument()

    doc.add_heading(conversation.title or "Conversation", level=1)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Espace : {space_name} — Exporte le {datetime.utcnow().strftime('%d/%m/%Y a %H:%M')} UTC"
    )
    meta_run.italic = True
    meta_run.font.size = Pt(9)

    for message in conversation.messages:
        if message.role not in ("user", "assistant"):
            continue
        label = "Question" if message.role == "user" else "Reponse"
        heading = doc.add_paragraph()
        run = heading.add_run(label)
        run.bold = True
        run.font.size = Pt(12)

        doc.add_paragraph(message.content)

        sources = _load_sources(message)
        if sources:
            for i, source in enumerate(sources, start=1):
                src_p = doc.add_paragraph(_source_line(source, i))
                for run in src_p.runs:
                    run.italic = True
                    run.font.size = Pt(9)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
